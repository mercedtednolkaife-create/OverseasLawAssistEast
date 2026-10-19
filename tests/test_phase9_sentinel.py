"""Phase 9 单元测试套件：双擎推理流、安全哨兵与央国企合规体系验证
测试覆盖：
1. 确定性特征引信与三层安全哨兵触发 (Tier 1 & Tier 2)
2. 建设工程专业审查算子：TKDN 本地化率、工期延误 (EOT) 与投建营接口传导
3. 质量门禁 Critic V12 央国企国家安全与涉外合规一票否决
4. 东盟建设工程专属入口检查清单 (calibrate_asean_entry_checks)
5. Word 双模板导出：东盟查明意见书与《央企境外工程专项合规指引手册 (DOCX)》
6. S1~S3 节点东盟工程业态自适应抽取、预检与争点拆解
"""

import pytest
from io import BytesIO
from docx import Document
from lawagent.core.types import (
    CaseState,
    CaseMeta,
    Fact,
    Claim,
    SecurityRiskLevel,
    SecurityAlertItem,
    CheckItem,
    SavedInsight
)
from lawagent.flow.security_sentinel import (
    SecuritySentinel,
    evaluate_tkdn_compliance,
    evaluate_eot_delay_claim,
    evaluate_interface_transmission,
    check_critic_v12_security_gate
)
from lawagent.flow.checks import generate_entry_checks, calibrate_asean_entry_checks
from lawagent.flow.critic import run_critic_checks
from lawagent.flow.nodes import (
    s1_extract_node,
    s2_precheck_node,
    s3_decompose_node,
    researcher_node,
    drafter_node
)
from lawagent.api.export import generate_case_docx, generate_compliance_handbook_docx


def test_security_sentinel_triggers():
    """测试 Tier 1 确定性引信与 Tier 2 安全哨兵触发完整性"""
    # 1. 隐性征收与单方废约触发
    state_expropriation = CaseState(
        case_id="TEST-SEC-01",
        snapshot_id="SNAP-01",
        facts_zh="东道国地方政府换届后单方废约，强行终止特许经营权并追溯下调FIT电价补贴，拒付重置成本补偿。",
        claims_zh="请求裁定政府单方废约违法并依据中菲双边投资协定主张公允补偿。",
        meta=CaseMeta(jurisdiction="asean_ph")
    )
    alerts = SecuritySentinel.evaluate(state_expropriation)
    assert len(alerts) > 0
    categories = [a.category for a in alerts]
    assert "creeping_expropriation" in categories
    
    crit_alert = next(a for a in alerts if a.category == "creeping_expropriation")
    assert crit_alert.risk_level == SecurityRiskLevel.CRITICAL
    assert "BIT_CHINA_ASEAN:Art6" in crit_alert.prc_or_treaty_basis
    assert crit_alert.sinosure_preservation_guide is not None
    assert "30 个自然日" in crit_alert.sinosure_preservation_guide or "30" in crit_alert.sinosure_preservation_guide

    # 2. 关键基础设施国安审查
    state_infra = CaseState(
        case_id="TEST-SEC-02",
        snapshot_id="SNAP-02",
        facts_zh="项目涉及印尼某特高压输变电及深水港核心基础设施工程总承包。",
        claims_zh="办理属地特许经营准入手续。",
        meta=CaseMeta(jurisdiction="asean_id")
    )
    alerts_infra = SecuritySentinel.evaluate(state_infra)
    assert any(a.category == "national_security_review" for a in alerts_infra)

    # 3. 数据安全法第 36 条跨境调证
    state_data = CaseState(
        case_id="TEST-SEC-03",
        snapshot_id="SNAP-03",
        facts_zh="境外仲裁庭发布强制命令，要求中资总承包商直接向境外调证提供国内服务器存储的工程地质测绘图纸及施工深化数据。",
        claims_zh="应对境外调证命令。",
        meta=CaseMeta(jurisdiction="asean_sg")
    )
    alerts_data = SecuritySentinel.evaluate(state_data)
    assert any(a.category == "cross_border_data_art36" for a in alerts_data)
    art36_alert = next(a for a in alerts_data if a.category == "cross_border_data_art36")
    assert art36_alert.risk_level == SecurityRiskLevel.CRITICAL
    assert "PRC_DATA_SEC:Art36" in art36_alert.prc_or_treaty_basis


def test_specialized_engineering_evaluators():
    """测试 TKDN 本地化率、工期延误 (EOT) 与投建营接口传导审查算子"""
    facts = (
        "爪哇光伏电站 EPC 工程中，业主迟延移交原住民习惯法土地 Hak Ulayat 达 180 天；"
        "同时因光伏逆变器本地化率 TKDN 仅达 28%（法定 40%），业主威胁扣罚违约金并拒绝顺延工期。"
    )
    
    # 1. TKDN 评估
    tkdn_eval = evaluate_tkdn_compliance(facts, jurisdiction="asean_id")
    assert tkdn_eval["applicable"] is True
    assert tkdn_eval["statutory_basis"] == "ID_MEMR:Reg11_2024"
    assert len(tkdn_eval["components_breakdown"]) >= 3
    assert "MEMR Reg 11/2024" in tkdn_eval["recommendation"] or "豁免" in tkdn_eval["recommendation"]

    # 2. EOT 工期延误评估
    eot_eval = evaluate_eot_delay_claim(facts, contract_type="FIDIC_SILVER", jurisdiction="asean_id")
    assert eot_eval["applicable"] is True
    assert "FIDIC_SILVER:Clause8.4" in eot_eval["fidic_clauses"]
    assert "FIDIC_SILVER:Clause20.1" in eot_eval["fidic_clauses"]
    assert len(eot_eval["elements_review"]) == 4

    # 3. 投建营接口传导审查 (马来西亚 CIPAA 35 条背靠背无效)
    interface_eval = evaluate_interface_transmission(
        facts_text="总承包商因业主未付款暂缓向本地分包商支付进度款，面临快速裁决风险。",
        archetypes=["EPC_CONTRACTING", "PPP_CONCESSION"],
        jurisdiction="asean_my"
    )
    assert interface_eval["applicable"] is True
    assert any("CIPAA 2012 第 35 条" in r["legal_rule"] for r in interface_eval["interface_risks"])


def test_critic_v12_security_gate():
    """测试质量门禁 Critic V12 央国企国家安全与涉外合规门禁"""
    # 情况 1: 未触发任何安全警报，V12 常规通过
    clean_state = CaseState(
        case_id="CLEAN-01",
        snapshot_id="SNAP-CLEAN",
        facts_zh="常规工程合同履约中未涉及任何国家安全与隐性征收问题"
    )
    passed, status, reason = check_critic_v12_security_gate(clean_state)
    assert passed is True
    assert status == "pass"

    # 情况 2: 触发 CRITICAL 警报但正文包含合规对冲及中信保保全，V12 通过
    critical_alert = SecurityAlertItem(
        risk_level=SecurityRiskLevel.CRITICAL,
        category="creeping_expropriation",
        title_zh="东道国政府单方废约与隐性征收",
        mitigation_actions=["依中信保 30 日出险保全程序提交报损通知书"]
    )
    mitigated_state = CaseState(
        case_id="MITIGATED-01",
        snapshot_id="SNAP-MIT",
        facts_zh="东道国单方废约特许经营权",
        security_alerts=[critical_alert],
        saved_insights=[
            SavedInsight(
                insight_id="INS-01",
                case_id="MITIGATED-01",
                category="memo",
                title="中信保保全与BIT抗辩",
                content="项目部已严格落实中信保 30 日出险报损留痕并准备 BIT 投资仲裁抗辩"
            )
        ]
    )
    passed_m, status_m, reason_m = check_critic_v12_security_gate(mitigated_state)
    assert passed_m is True
    assert status_m == "pass"

    # 校验 run_critic_checks 整合包含 V12
    checks, dashboard = run_critic_checks(mitigated_state)
    v12_item = next((c for c in checks if c.code == "V12"), None)
    assert v12_item is not None
    assert v12_item.status == "pass"


def test_asean_entry_checks_calibration():
    """测试东盟专属入口检查清单生成逻辑"""
    meta_id = CaseMeta(
        jurisdiction="asean_id",
        choice_of_law="印尼法律及 FIDIC 银皮书",
        contract_type="EPC_CONTRACTING"
    )
    checks = generate_entry_checks(meta_id)
    assert len(checks) == 5
    check_names = [c.name for c in checks]
    assert any("SEMA 3/2023" in n for n in check_names)
    assert any("BUJKA" in n or "施工牌照" in n for n in check_names)
    assert any("国资委 35 号令" in n for n in check_names)

    meta_my = CaseMeta(
        jurisdiction="asean_my",
        choice_of_law="马来西亚法律",
        contract_type="EPC_CONTRACTING"
    )
    checks_my = generate_entry_checks(meta_my)
    assert any("CIPAA 2012 第 35 条" in c.detail for c in checks_my)


def test_export_case_and_compliance_handbook():
    """测试 Word 双模板导出：查明意见书与《专项合规指引手册 (DOCX)》"""
    alert = SecurityAlertItem(
        risk_level=SecurityRiskLevel.CRITICAL,
        category="creeping_expropriation",
        title_zh="特许权单方终止与隐性征收",
        fact_anchor_ids=["F1"],
        host_law_basis=["ID_UU:2_2017:Art57"],
        prc_or_treaty_basis=["BIT_CHINA_ASEAN:Art6"],
        sinosure_preservation_guide="须在 30 个自然日内向中国信保提交可能损失通知书并保全代位求偿权。",
        mandatory_human_checklist=["核实中信保保单缴费记录", "严禁签署放弃索赔纪要"],
        mitigation_actions=["启动 BIT 条约磋商", "申请 SIAC 紧急仲裁员临时禁令"]
    )
    case_state = CaseState(
        case_id="ASEAN-DOCX-01",
        snapshot_id="SNAP-DOCX",
        facts_zh="雅万电站项目业主迟延交地，单方停工并威胁索兑履约保函。",
        meta=CaseMeta(
            jurisdiction="asean_id",
            choice_of_law="印尼法及 FIDIC 银皮书",
            contract_type="国际建设工程EPC总承包合同",
            project_archetypes=["EPC_CONTRACTING", "JV_CONSORTIUM"]
        ),
        security_alerts=[alert]
    )

    # 1. 导出案件查明咨询报告
    case_doc_stream = generate_case_docx(case_state)
    assert isinstance(case_doc_stream, BytesIO)
    assert case_doc_stream.getbuffer().nbytes > 10000
    doc = Document(case_doc_stream)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "东盟国家建设工程与海外投资跨境查明与合规咨询报告" in full_text
    assert "国家安全、海外投资经济安全与反隐性征收专项预警" in full_text
    assert "中国信保 (Sinosure) 30 日出险报损保全指引" in full_text

    # 2. 导出《专项合规指引手册 (DOCX)》
    handbook_stream = generate_compliance_handbook_docx(
        jurisdiction="asean_id",
        archetypes=["EPC_CONTRACTING", "PPP_CONCESSION"]
    )
    assert isinstance(handbook_stream, BytesIO)
    assert handbook_stream.getbuffer().nbytes > 10000
    hb_doc = Document(handbook_stream)
    hb_text = "\n".join([p.text for p in hb_doc.paragraphs])
    assert "专项合规指引手册 (SOP)" in hb_text
    assert "阶段 I：前期跟踪与立项尽调" in hb_text
    assert "国资委“三道防线”角色实操审查清单" in hb_text
    assert "反习惯法征拆迟延免责条款" in hb_text
    assert "原住民习惯法土地" in hb_text


@pytest.mark.asyncio
async def test_asean_flow_s1_s2_s3_nodes():
    """测试 S1~S3 节点对东盟工程建设业态的自适应抽取与争点拆解"""
    init_state = CaseState(
        case_id="FLOW-ASEAN-01",
        snapshot_id="SNAP-FLOW",
        facts_zh="西爪哇太阳能电站 EPC 项目中，印尼业主迟延移交习惯法土地 180 天，总承包商依 FIDIC 8.4 发函索赔工期，业主以纯英文备忘录无效及逆变器本地化率不足为由拒付并威胁索兑保函。",
        claims_zh="1. 索赔工期顺延 180 天；2. 依据 SEMA 3/2023 确认英文协议有效；3. 申请保函止付禁令。",
        meta=CaseMeta(jurisdiction="asean_id", lifecycle_stage="execution_dispute")
    )

    # S1 Extract
    s1_res = await s1_extract_node(init_state)
    assert s1_res["meta"].contract_type == "国际建设工程EPC/特许经营合同"
    assert "中资总承包商/投资方" in s1_res["meta"].parties
    for k, v in s1_res.items():
        if hasattr(init_state, k):
            setattr(init_state, k, v)

    # S2 Precheck
    s2_res = await s2_precheck_node(init_state)
    assert "security_alerts" in s2_res
    assert s2_res["status"] == "running"  # 东盟工程非国际货物买卖，不暂停 CISG 闸门
    for k, v in s2_res.items():
        if hasattr(init_state, k):
            setattr(init_state, k, v)

    # S3 Decompose
    s3_res = await s3_decompose_node(init_state)
    assert len(s3_res["issues"]) == 3
    assert len(s3_res["questions"]) >= 1
    assert any("FIDIC_SILVER:Clause8.4" in q.must_norms for q in s3_res["questions"])
    assert any("ID_SEMA:3_2023" in q.must_norms for q in s3_res["questions"])
