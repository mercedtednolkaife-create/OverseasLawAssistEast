"""东盟涉外建设工程与海外投资专项全链路验收测试套件 (ASEAN Architecture Test Suite)
对齐《东盟国家迁移方案.md》第八节与第九节 Phase 10 质量验收规范：
1. 四大金标准示范案 (C101~C104) 数据完整性、法条锚点验真与 0 幻觉率核验；
2. 投前合规 (Mode 1) 与履约争议 (Mode 2) 生命周期阶段双向流转与契约继承；
3. 最高人民法院查明意见书与《项目专项合规指引手册》Word (DOCX) 双模板导出闭环；
4. 国家安全哨兵、隐性征收、中信保 30 日出险报损与 Critic V12 门禁 100% 熔断对冲；
5. 3-Layer Legal GraphRAG 拓扑边与端到端工作流自适应推进。
"""

import os
import pytest
from io import BytesIO
from docx import Document

from lawagent.core.types import (
    CaseState, CaseMeta, Fact, Claim, Material, LegalQuestion, Issue, EntryCheckItem,
    ElementState, Evidence, Finding, Statement, CheckItem, DashboardMetrics,
    LifecycleStage, ProjectArchetype, SecurityRiskLevel, SecurityAlertItem
)
from lawagent.store.db import StateDBManager
from lawagent.flow.security_sentinel import (
    SecuritySentinel,
    evaluate_tkdn_compliance,
    evaluate_eot_delay_claim,
    evaluate_interface_transmission,
    check_critic_v12_security_gate
)
from lawagent.flow.critic import run_critic_checks
from lawagent.flow.checks import generate_entry_checks
from lawagent.flow.nodes import (
    s1_extract_node,
    s2_precheck_node,
    s3_decompose_node
)
from lawagent.api.export import generate_case_docx, generate_compliance_handbook_docx
from scripts.seed_asean_cases import (
    create_c101_case,
    create_c102_case,
    create_c103_case,
    create_c104_case,
    seed_asean_demo_cases
)


@pytest.fixture(scope="module")
def seeded_db():
    """初始化数据库并播种东盟四大金标准示范案"""
    db = StateDBManager("data/state/state.sqlite")
    seed_asean_demo_cases(db, overwrite=True)
    return db


def test_asean_demonstration_cases_integrity(seeded_db):
    """测试 1: 验证 C101~C104 四大东盟标杆示范案的数据完整性、2026法源真实性与零幻觉率"""
    for cid in ["C101", "C102", "C103", "C104"]:
        state = seeded_db.get_case(cid)
        assert state is not None, f"Case {cid} must exist in state.sqlite"
        assert state.case_id == cid
        assert state.status == "completed"
        assert state.snapshot_id == "20261018-asean-v1.0"
        assert len(state.facts) >= 3
        assert len(state.claims) >= 3
        assert len(state.materials) >= 2
        assert len(state.issues) >= 2
        assert len(state.evidence) >= 2
        assert len(state.findings) >= 1
        assert state.dashboard.content_hash != ""
        assert state.meta.jurisdiction.startswith("asean_")

    # 1. C101 (印尼): 习惯法土地、SEMA 3/2023、MEMR Reg 11/2024 与 FIDIC 8.4/20.1
    c101 = seeded_db.get_case("C101")
    assert c101.meta.jurisdiction == "asean_id"
    assert c101.meta.lifecycle_stage == LifecycleStage.EXECUTION_DISPUTE
    assert ProjectArchetype.EPC_CONTRACTING in c101.meta.project_archetypes
    assert ProjectArchetype.JV_CONSORTIUM in c101.meta.project_archetypes
    ref_ids_101 = [ev.ref_id for ev in c101.evidence.values()]
    assert "ID_UU:2_2017:Art57" in ref_ids_101
    assert "ID_SEMA:3_2023" in ref_ids_101
    assert "ID_MEMR:Reg11_2024" in ref_ids_101
    assert "FIDIC_SILVER:Clause8.4" in ref_ids_101
    assert any(a.category == "resource_nationalism_tkdn" for a in c101.security_alerts)

    # 2. C102 (越南): 2025新建设法 86.2 条预定违约金 12% 封顶与不可抗力
    c102 = seeded_db.get_case("C102")
    assert c102.meta.jurisdiction == "asean_vn"
    ref_ids_102 = [ev.ref_id for ev in c102.evidence.values()]
    assert "VN_LAW:135_2025_QH15:Art86" in ref_ids_102
    assert "VN_DECREE:210_2026_ND_CP" in ref_ids_102
    assert any(a.category == "creeping_expropriation" for a in c102.security_alerts)

    # 3. C103 (马来西亚/新加坡): CIPAA 35 条反背靠背与新加坡 BS Mount Sophia 保函显失公平止付
    c103 = seeded_db.get_case("C103")
    assert c103.meta.jurisdiction == "asean_my"
    ref_ids_103 = [ev.ref_id for ev in c103.evidence.values()]
    assert "MY_CIPAA:Sec35" in ref_ids_103
    assert "MY_CIPAA:Sec13" in ref_ids_103
    assert "SG_CASE:BS_Mount_Sophia" in ref_ids_103

    # 4. C104 (菲律宾): PPP法典 RA 11966、最高院 PCAB v. Manila Water 废除外资股权限制与中菲 BIT
    c104 = seeded_db.get_case("C104")
    assert c104.meta.jurisdiction == "asean_ph"
    ref_ids_104 = [ev.ref_id for ev in c104.evidence.values()]
    assert "PH_RA:11966:PPP_Code" in ref_ids_104
    assert "PH_SC:GR217590_ManilaWater" in ref_ids_104
    assert "BIT_CHINA_ASEAN:Art6" in ref_ids_104
    assert any(a.sinosure_preservation_guide is not None for a in c104.security_alerts)


def test_stage_transition_roundtrip(seeded_db):
    """测试 2: 验证投前审查与履约争议生命周期阶段双向流转及元数据继承机制"""
    # 以 C102 (PRE_INVESTMENT) 为蓝本派生流转至 EXECUTION_DISPUTE
    c102 = seeded_db.get_case("C102")
    assert c102.meta.lifecycle_stage == LifecycleStage.PRE_INVESTMENT

    derived_case = CaseState(
        case_id="C102-EXEC",
        parent_case_id="C102",
        title="越南海上风电股权并购项目-履约索赔与保函止付派生案",
        snapshot_id=c102.snapshot_id,
        mode="full",
        status="running",
        facts_zh=c102.facts_zh + " 经双方投前交割后，业主已正式向开证行提交没收履约保函索赔信函。",
        claims_zh=c102.claims_zh,
        meta=CaseMeta(
            title="越南海上风电-履约争议阶段",
            jurisdiction=c102.meta.jurisdiction,
            contract_type=c102.meta.contract_type,
            choice_of_law=c102.meta.choice_of_law,
            lifecycle_stage=LifecycleStage.EXECUTION_DISPUTE,
            project_archetypes=c102.meta.project_archetypes,
            parent_case_id="C102"
        ),
        security_alerts=c102.security_alerts
    )
    seeded_db.save_case(derived_case, title=derived_case.title)

    saved_derived = seeded_db.get_case("C102-EXEC")
    assert saved_derived is not None
    assert saved_derived.parent_case_id == "C102"
    assert saved_derived.meta.lifecycle_stage == LifecycleStage.EXECUTION_DISPUTE
    assert saved_derived.meta.jurisdiction == "asean_vn"
    assert ProjectArchetype.CROSS_BORDER_MA in saved_derived.meta.project_archetypes
    assert ProjectArchetype.EPC_CONTRACTING in saved_derived.meta.project_archetypes

    # 清理测试派生案
    seeded_db.delete_case("C102-EXEC")


def test_asean_dual_export_docx_validation(seeded_db):
    """测试 3: 验证 Word 双模板导出：查明意见书与《项目专项合规指引手册 (DOCX)》"""
    # 1. 验证 C101 (印尼) 查明意见书 DOCX
    c101 = seeded_db.get_case("C101")
    c101_stream = generate_case_docx(c101)
    assert isinstance(c101_stream, BytesIO)
    assert c101_stream.getbuffer().nbytes > 12000
    doc_101 = Document(c101_stream)
    text_101 = "\n".join([p.text for p in doc_101.paragraphs] + [c.text for t in doc_101.tables for r in t.rows for c in r.cells])
    assert "东盟国家建设工程与海外投资跨境查明与合规咨询报告" in text_101
    assert "国家安全、海外投资经济安全与反隐性征收专项预警" in text_101
    assert "印尼 TKDN 本地化率合规风险" in text_101
    assert "FIDIC 银皮书" in text_101
    assert "中国信保 (Sinosure) 30 日出险报损保全指引" in text_101

    # 2. 验证 C104 (菲律宾) 查明意见书包含信保 30 日出险指引
    c104 = seeded_db.get_case("C104")
    c104_stream = generate_case_docx(c104)
    doc_104 = Document(c104_stream)
    text_104 = "\n".join([p.text for p in doc_104.paragraphs] + [c.text for t in doc_104.tables for r in t.rows for c in r.cells])
    assert "中国信保 (Sinosure) 30 日出险报损保全指引" in text_104
    assert "RA 11966" in text_104
    assert "PCAB v. Manila Water" in text_104

    # 3. 验证《涉外工程专项合规指引手册 (DOCX)》导出
    handbook_stream = generate_compliance_handbook_docx(
        jurisdiction="asean_id",
        archetypes=["EPC_CONTRACTING", "PPP_CONCESSION"]
    )
    assert isinstance(handbook_stream, BytesIO)
    assert handbook_stream.getbuffer().nbytes > 12000
    hb_doc = Document(handbook_stream)
    hb_text = "\n".join([p.text for p in hb_doc.paragraphs] + [c.text for t in hb_doc.tables for r in t.rows for c in r.cells])
    assert "专项合规指引手册 (SOP)" in hb_text
    assert "建设工程与跨境投资" in hb_text
    assert "阶段 I：前期跟踪与立项尽调" in hb_text
    assert "阶段 III：合同谈判与签署红线" in hb_text
    assert "阶段 V：竣工移交、争议解决与退出清算" in hb_text
    assert "【第一道防线】海外市场部 / 项目现场经理审查要点" in hb_text
    assert "【第二道防线】涉外法务总监 / 专职合规官审查要点" in hb_text
    assert "反习惯法征拆迟延免责条款" in hb_text
    assert "见索即付保函索兑前置限制与止付约定" in hb_text


def test_critic_v12_security_gate_full_closure(seeded_db):
    """测试 4: 验证四大标杆案例 Critic V12 央国企国家安全门禁 100% 通过且具备有效对冲"""
    for cid in ["C101", "C102", "C103", "C104"]:
        state = seeded_db.get_case(cid)
        # 验证单独门禁校验
        passed, status, reason = check_critic_v12_security_gate(state)
        assert passed is True, f"Case {cid} must pass Critic V12 gate"
        assert status == "pass"

        # 验证与 run_critic_checks 整合执行
        checks, dashboard = run_critic_checks(state)
        v12 = next((c for c in checks if c.code == "V12"), None)
        assert v12 is not None, f"V12 check item must be present for ASEAN Case {cid}"
        assert v12.status == "pass"


@pytest.mark.asyncio
async def test_end_to_end_asean_workflow_simulation():
    """测试 5: 端到端东盟建设工程工作流模拟 (S1 -> S2 -> S3) 验证 CISG 闸门免暂停与争点拆解"""
    sim_state = CaseState(
        case_id="SIM-ASEAN-01",
        snapshot_id="20261018-asean-v1.0",
        facts_zh="西爪哇太阳能电站 EPC 工程中，业主迟延移交原住民习惯法土地达 180 天，总承包商依 FIDIC 8.4 发函索赔工期，业主以纯英文备忘录无效及逆变器本地化率不足为由拒付并威胁索兑保函。",
        claims_zh="1. 索赔工期顺延 180 天；2. 依据 SEMA 3/2023 确认英文协议有效；3. 申请保函止付禁令。",
        meta=CaseMeta(jurisdiction="asean_id", lifecycle_stage=LifecycleStage.EXECUTION_DISPUTE)
    )

    # 1. S1: 事实与业态自适应抽取
    s1_res = await s1_extract_node(sim_state)
    assert s1_res["meta"].contract_type == "国际建设工程EPC/特许经营合同"
    assert "中资总承包商/投资方" in s1_res["meta"].parties
    for k, v in s1_res.items():
        if hasattr(sim_state, k):
            setattr(sim_state, k, v)

    # 2. S2: 前检与安全哨兵触发 (验证不因货物买卖而暂停 CISG 闸门)
    s2_res = await s2_precheck_node(sim_state)
    assert s2_res["status"] == "running"  # 核心：东盟基建非 CISG 货物买卖，不暂停
    assert len(s2_res["security_alerts"]) >= 1
    assert any(a.category == "national_security_review" for a in s2_res["security_alerts"])
    for k, v in s2_res.items():
        if hasattr(sim_state, k):
            setattr(sim_state, k, v)

    # 3. S3: 东盟建设工程专属争点与构成要件拆解
    s3_res = await s3_decompose_node(sim_state)
    assert len(s3_res["issues"]) == 3
    assert len(s3_res["questions"]) >= 1
    q0 = s3_res["questions"][0]
    assert any("FIDIC_SILVER:Clause8.4" in n for n in q0.must_norms)
    assert any("ID_SEMA:3_2023" in n for n in q0.must_norms)
    assert any("MY_CIPAA:Sec35" in n for n in q0.must_norms)


def test_sinosure_compliance_operator():
    """核验中信保 30 日出险报损时效审查算子"""
    from lawagent.flow.security_sentinel import evaluate_sinosure_compliance

    # Case A: 12天报损，合规安全
    res_safe = evaluate_sinosure_compliance("中信保海外投资险保单出险报损", engineering_params={"sinosure_notice_days": 12})
    assert res_safe["applicable"] is True
    assert res_safe["is_compliant"] is True
    assert res_safe["is_overdue"] is False
    assert res_safe["days_remaining"] == 18
    assert res_safe["status"] == "safe"

    # Case B: 38天报损，逾期预警
    res_overdue = evaluate_sinosure_compliance("中信保出险报损通知", engineering_params={"sinosure_notice_days": 38})
    assert res_overdue["applicable"] is True
    assert res_overdue["is_compliant"] is False
    assert res_overdue["is_overdue"] is True
    assert res_overdue["days_remaining"] == 0
    assert res_overdue["status"] == "caution"


def test_ph_and_th_market_entry_operators():
    """核验菲律宾 PCAB 施工资质与泰国 FBA 外资商业法审查算子"""
    from lawagent.flow.security_sentinel import evaluate_ph_contractor_license, evaluate_th_foreign_business_act

    # 菲律宾 PCAB 资质测试
    res_ph_reg = evaluate_ph_contractor_license("菲律宾马尼拉大桥工程", foreign_equity_pct=35.0)
    assert res_ph_reg["applicable"] is True
    assert res_ph_reg["is_regular_eligible"] is True
    assert "Regular License" in res_ph_reg["recommended_license"]

    res_ph_spec = evaluate_ph_contractor_license("菲律宾公工部 PPP 项目", foreign_equity_pct=60.0)
    assert res_ph_spec["applicable"] is True
    assert res_ph_spec["is_regular_eligible"] is False
    assert "Special License" in res_ph_spec["recommended_license"]
    assert "Manila_Water" in res_ph_spec["legal_basis"]
    assert "Manila Water" in res_ph_spec["counter_measure"]

    # 泰国 FBA 外资商业法测试
    res_th_ok = evaluate_th_foreign_business_act("泰国曼谷电站 EPC 项目", foreign_equity_pct=45.0)
    assert res_th_ok["applicable"] is True
    assert res_th_ok["is_fba_restricted"] is False

    res_th_restricted = evaluate_th_foreign_business_act("泰国曼谷高架桥施工项目", foreign_equity_pct=70.0)
    assert res_th_restricted["applicable"] is True
    assert res_th_restricted["is_fba_restricted"] is True
    assert "BOI" in res_th_restricted["exemption_route"]


def test_lex_loci_solutionis_decoupling_override():
    """核验合同约定准据法与工程所在国强行法脱钩穿透审查门禁"""
    # 脱钩场景：项目在印尼，约定英国法/新加坡法
    meta_decoupled = CaseMeta(
        jurisdiction="asean_id",
        project_host_country="asean_id",
        governing_law_jurisdiction="英国普通法 (English Law)"
    )
    checks_decoupled = generate_entry_checks(meta_decoupled)
    override_items = [c for c in checks_decoupled if "脱钩" in c.name or "Lex Loci" in c.name]
    assert len(override_items) == 1
    assert "不可约定排除" in override_items[0].detail

    # 未脱钩场景：项目在印尼，约定印尼实体法
    meta_same = CaseMeta(
        jurisdiction="asean_id",
        project_host_country="asean_id",
        governing_law_jurisdiction="asean_id"
    )
    checks_same = generate_entry_checks(meta_same)
    assert not any("脱钩" in c.name for c in checks_same)


def test_local_counsel_expert_roster_coverage(seeded_db):
    """核验东盟六国 Local Counsel 与 DAAB 专家库名录完整性"""
    experts = seeded_db.list_experts()
    assert len(experts) >= 6
    countries_represented = set()
    for exp in experts:
        aff = exp["affiliation"] + " " + exp["name"] + " " + exp["title"]
        if "印尼" in aff or "BANI" in aff or "Wijaya" in aff:
            countries_represented.add("ID")
        if "新加坡" in aff or "SIAC" in aff or "Kevin Tan" in aff:
            countries_represented.add("SG")
        if "马来西亚" in aff or "AIAC" in aff or "Lim Chong Beng" in aff:
            countries_represented.add("MY")
        if "越南" in aff or "VIAC" in aff or "Nguyen" in aff:
            countries_represented.add("VN")
        if "泰国" in aff or "TAI" in aff or "Somchai" in aff:
            countries_represented.add("TH")
        if "菲律宾" in aff or "CIAC" in aff or "Santos" in aff:
            countries_represented.add("PH")

    assert {"ID", "SG", "MY", "VN", "TH", "PH"}.issubset(countries_represented)


@pytest.mark.asyncio
async def test_async_verify_foreign_materials():
    """核验外国法证据材料异步线程池校验与红线扫描"""
    from lawagent.api.routes.verify import verify_foreign_materials, VerifyRequest
    from lawagent.store.snapshots import SnapshotManager

    sm = SnapshotManager()
    sample_text = (
        "Clause 14.8: Payments shall be pay when paid. "
        "FIDIC Clause 20.1 requires notice within 28 days. "
        "Under BS Mount Sophia, injunction was granted."
    )
    req = VerifyRequest(raw_text=sample_text, mode="CONTRACT_REDLINE_AUDIT")
    res = await verify_foreign_materials(req, snapshot_mgr=sm)

    assert res.raw_text == sample_text
    assert len(res.redline_items) >= 1
    assert any("背靠背" in item.violated_rule_title for item in res.redline_items)
    assert len(res.items) >= 1
    assert res.summary["total"] >= 1

