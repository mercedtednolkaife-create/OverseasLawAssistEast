"""Word 报告与合规手册导出模块 (DOCX Report & Compliance Handbook Generator)
严格满足《最高人民法院关于适用〈涉外民事关系法律适用法〉若干问题的解释（二）》第三条，
以及国务院国资委《中央企业境外投资监督管理办法》（35号令）合规交底要求：
1. 案件查明报告 (generate_case_docx)：
   - A包 域外法规范查明材料包 (双语/三语对照、条文指纹、效力说明)
   - B包 事实涵摄与咨询准备包 (要件矩阵、律师咨询清单、先例比对)
   - 国家安全与海外投资经济安全哨兵专项预警包 (中信保30日报损、国资监管、BIT条约)
2. 项目专项合规指引手册 (generate_compliance_handbook_docx)：
   - 全生命周期五阶段 SOP
   - 国资委“三道防线”实操审查清单
   - 关键工程合同示范反制条款 (Model Counter-Clauses)
"""

import os
from io import BytesIO
from datetime import datetime
from typing import List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from lawagent.core.types import CaseState, SecurityRiskLevel


def _set_cell_background(cell, fill_hex: str):
    """设置单元格背景底色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def generate_case_docx(state: CaseState) -> BytesIO:
    """基于 CaseState 生成结构完备的专业法律查明咨询报告"""
    doc = Document()


    # 1. 设置主标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_title_text = "东盟国家建设工程与海外投资跨境查明与合规咨询报告"
    sub_title_text = "ASEAN Infrastructure & Cross-Border Investment Legal Ascertainment & Advisory Dossier"

    title_run = title_p.add_run(main_title_text)
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(sub_title_text)
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # 2. 案件元数据摘要表
    h1 = doc.add_heading("一、查明基本信息与案件元数据", level=1)
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    default_parties = "中资总承包商/投资方 与 东道国业主/政府方"
    default_choice = f"东盟东道国法 ({state.meta.jurisdiction}) 及 FIDIC 标准合同"
    default_snapshot = "司法部法通网 / 2026现行东盟成文法库"

    labels = [
        ("案件查明编号", state.case_id),
        ("合同性质与当事人", f"{state.meta.contract_type or '国际建设工程EPC/特许经营合同'} | {', '.join(state.meta.parties) or default_parties}"),
        ("选法条款与公约适用", f"准据法：{state.meta.choice_of_law or default_choice} | 国际公约/条约：依中国-东盟双边投资协定 (BIT) 与东道国强行法"),
        ("关键时点与版本核验", f"合同日：{state.meta.contract_date or '未载明'} | 事实日：{state.meta.event_date or '未载明'} | 时点警示：【{state.timepoint.get('level', 'yellow').upper()}态】"),
        ("官方快照数据指纹", f"不可变快照：{state.snapshot_id} ({default_snapshot})")
    ]
    archetypes_str = ", ".join([a.value if hasattr(a, 'value') else str(a) for a in (state.meta.project_archetypes or [])]) or "EPC_CONTRACTING"
    stage_str = "履约索赔与争议解决" if (state.meta.lifecycle_stage == "execution_dispute" or getattr(state.meta.lifecycle_stage, 'value', '') == "execution_dispute") else "投前并购准入与签约审查"
    labels.append(("工程业态与生命周期", f"生命周期：{stage_str} | 业态组合：{archetypes_str} | 安全预警：{len(state.security_alerts or [])}项触发"))

    for idx, (lbl, val) in enumerate(labels):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        cell_lbl.text = lbl
        cell_val.text = val
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        _set_cell_background(cell_lbl, "F1F5F9")

    doc.add_paragraph()

    # 3. 第一部分：A包 域外法规范查明材料包
    part_a_title = "二、第一部分：东盟现行有效工程法规与 FIDIC 标准合同材料包 (Part A)"
    p_a_intro_text = "依据司法部‘法通’网权威译本、东盟东道国2026现行有效成文法及 FIDIC 国际工程标准合同，系统严格定位并提取支撑本案争议焦点之法定权利基础及构成要件条款如下："
    col2_label = "官方语言原文 (Statute / FIDIC)"

    doc.add_heading(part_a_title, level=1)
    p_a_intro = doc.add_paragraph(p_a_intro_text)
    p_a_intro.runs[0].font.color.rgb = RGBColor(51, 65, 85)

    if state.evidence:
        norm_evs = [ev for ev in state.evidence.values() if ev.kind == "norm"]
        if norm_evs:
            doc.add_heading("2.1 核心法定依据多语对照清单", level=2)
            ev_table = doc.add_table(rows=len(norm_evs) + 1, cols=4)
            ev_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            headers = ["编号", "规范条号", col2_label, "中文权威释义与术语对照"]
            for col_idx, h in enumerate(headers):
                cell = ev_table.rows[0].cells[col_idx]
                cell.text = h
                cell.paragraphs[0].runs[0].font.bold = True
                _set_cell_background(cell, "E2E8F0")

            ev_table.columns[0].width = Inches(0.8)
            ev_table.columns[1].width = Inches(1.3)
            ev_table.columns[2].width = Inches(2.2)
            ev_table.columns[3].width = Inches(2.2)

            for r_idx, ev in enumerate(norm_evs, 1):
                row = ev_table.rows[r_idx]
                row.cells[0].text = ev.eid
                row.cells[1].text = ev.ref_id
                row.cells[2].text = ev.text_de[:350] + ("..." if len(ev.text_de) > 350 else "")
                trans = state.translations.get(ev.eid, "中文释义待校验")
                row.cells[3].text = trans[:350] + ("..." if len(trans) > 350 else "")

    doc.add_paragraph()

    # 4. 第二部分：B包 事实涵摄与律师咨询准备包
    doc.add_heading("三、第二部分：事实涵摄与涉外法律咨询准备包 (Part B)", level=1)
    target_lawyer = "涉外工程仲裁大律师与央企海外合规法务"
    p_b_intro = doc.add_paragraph(
        f"本部分聚焦案件事实与法律构成要件之深度涵摄分析，提炼事实举证状态、案情缺口并形成向{target_lawyer}咨询之专业问询清单，供仲裁庭审与出海合规谈判决策之用。"
    )

    if state.elements:
        doc.add_heading("3.1 构成要件事实匹配与涵摄状态表", level=2)
        elem_table = doc.add_table(rows=len(state.elements) + 1, cols=4)
        elem_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        e_headers = ["构成要件", "关联事实凭证", "满足状态", "事实涵摄与法律分析说明"]
        for col_idx, h in enumerate(e_headers):
            cell = elem_table.rows[0].cells[col_idx]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            _set_cell_background(cell, "E2E8F0")

        elem_table.columns[0].width = Inches(1.5)
        elem_table.columns[1].width = Inches(1.2)
        elem_table.columns[2].width = Inches(1.0)
        elem_table.columns[3].width = Inches(2.8)

        status_map = {
            "satisfied": "充分满足",
            "partial": "部分满足",
            "missing": "事实缺失",
            "disputed": "存有争议"
        }

        for r_idx, el in enumerate(state.elements, 1):
            row = elem_table.rows[r_idx]
            row.cells[0].text = el.name_zh
            row.cells[1].text = ", ".join(el.fact_links or ["待举证"])
            st_text = status_map.get(el.match_status or "missing", "待定")
            row.cells[2].text = st_text
            row.cells[3].text = el.match_note or "需进一步核实要件事实"

    doc.add_paragraph()

    # 咨询问题
    if state.counsel_questions:
        doc.add_heading(f"3.2 向{target_lawyer}咨询之关键问题清单 (Counsel Questions)", level=2)
        for q_idx, cq in enumerate(state.counsel_questions, 1):
            p_q = doc.add_paragraph()
            p_q.add_run(f"Q{q_idx}. ").bold = True
            p_q.add_run(cq)

    # 类似判例或示范案先例
    if state.similar_cases:
        precedent_title = "3.3 东盟工程仲裁与司法判例要旨精选"
        doc.add_heading(precedent_title, level=2)
        for sc in state.similar_cases:
            p_c = doc.add_paragraph()
            p_c.add_run(f"【{sc.court} {sc.az} ({sc.date})】\n").bold = True
            p_c.add_run(f"裁判要旨：{sc.court_view_zh}\n")
            p_c.add_run(f"裁判原文片段：{sc.court_view_quote_de[:250]}...\n").italic = True
            p_c.add_run(f"事实异同对比：{'; '.join(sc.similarity_points)}；差异点：{'; '.join(sc.difference_points)}")

    doc.add_paragraph()

    # 5. 国家安全与海外投资经济安全哨兵专项预警章节 (Security Sentinel Dossier)
    if state.security_alerts:
        doc.add_heading("四、国家安全、海外投资经济安全与反隐性征收专项预警 (Security Sentinel)", level=1)
        p_sec_intro = doc.add_paragraph(
            "针对涉外基建项目可能涉及的主权征收、特许权撤销、制裁除名及数据跨境等高危风险，系统三层刚性哨兵已触发专项防御指示："
        )
        p_sec_intro.runs[0].font.color.rgb = RGBColor(185, 28, 28)

        for a_idx, alert in enumerate(state.security_alerts, 1):
            level_str = "【🔴 致命风险 CRITICAL】" if alert.risk_level == SecurityRiskLevel.CRITICAL else "【🟠 高危风险 HIGH】" if alert.risk_level == SecurityRiskLevel.HIGH else "【🟡 中度关注 CAUTION】"
            doc.add_heading(f"4.{a_idx} {level_str} {alert.title_zh}", level=2)
            
            p_alert_meta = doc.add_paragraph()
            p_alert_meta.add_run("触发事实锚点：").bold = True
            p_alert_meta.add_run(f"{', '.join(alert.fact_anchor_ids) or '全案事实综合判定'}\n")
            p_alert_meta.add_run("东道国法条及双边条约依据：").bold = True
            p_alert_meta.add_run(f"{', '.join(alert.host_law_basis + alert.prc_or_treaty_basis) or '国际投资法及双边协定'}\n")

            if alert.sinosure_preservation_guide:
                p_sino = doc.add_paragraph()
                p_sino.add_run("中国信保 (Sinosure) 30 日出险报损保全指引：\n").bold = True
                p_sino.add_run(alert.sinosure_preservation_guide)

            if alert.mandatory_human_checklist:
                p_check = doc.add_paragraph()
                p_check.add_run("央国企法务强制人工核验清单：\n").bold = True
                for item in alert.mandatory_human_checklist:
                    p_check.add_run(f"  □ {item}\n")

            if alert.mitigation_actions:
                p_mit = doc.add_paragraph()
                p_mit.add_run("反制与对冲建议 (Mitigation Actions)：\n").bold = True
                for act in alert.mitigation_actions:
                    p_mit.add_run(f"  • {act}\n")

        doc.add_paragraph()

    # 6. 质量门禁与防线声明
    disclaimer_level = 5 if state.security_alerts else 4
    doc.add_heading(f"{'五' if state.security_alerts else '四'}、质量门禁与客观中立性声明", level=1)
    bar_qual = "东盟东道国或国际工程仲裁执业资格"
    p_disclaimer = doc.add_paragraph(
        f"【客观中立性防线声明与合规依据】本报告依据《最高人民法院关于适用〈涉外民事关系法律适用法〉若干问题的解释（二）》关于涉外法域查明之规范指引、国务院国资委《中央企业境外投资监督管理办法》（第35号令）合规要求，并结合司法部法通网不可变东盟法律快照生成。本报告仅提供客观规范条款、构成要件逻辑结构、风险哨兵提示与证据准备指引，坚决不作任何绝对胜诉率预测或保函绝对止付保证。所有实体权利与程序主张均应由具有{bar_qual}之涉外法律专业人员线下最终复核确认。"
    )
    p_disclaimer.runs[0].font.color.rgb = RGBColor(180, 83, 9)

    doc_stream = BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    return doc_stream


def generate_compliance_handbook_docx(
    jurisdiction: str = "asean_id",
    archetypes: Optional[List[str]] = None
) -> BytesIO:
    """生成《中央企业东盟建设工程与海外投资项目专项合规指引手册 (DOCX)》
    对标国资委 35 号令与东盟东道国 2026 现行有效法源
    """
    doc = Document()
    archetypes = archetypes or ["EPC_CONTRACTING", "PPP_CONCESSION"]
    archetypes_text = "/".join(archetypes)

    jur_names = {
        "asean_id": "印度尼西亚共和国 (Indonesia)",
        "asean_vn": "越南社会主义共和国 (Vietnam)",
        "asean_my": "马来西亚 (Malaysia)",
        "asean_sg": "新加坡共和国 (Singapore)",
        "asean_th": "泰王国 (Thailand)",
        "asean_ph": "菲律宾共和国 (Philippines)"
    }
    jur_label = jur_names.get(jurisdiction, "东盟重点国家")

    # 1. 封面与手册标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_title = f"{jur_label}建设工程与跨境投资"
    title_run = title_p.add_run(f"【央企境外合规交底】\n{main_title}\n专项合规指引手册 (SOP)")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"业态涵盖：{archetypes_text} | 编审标准：国资委35号令 / 2026现行法通网权威版本")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # 2. 前言
    doc.add_heading("一、 指引手册定位与使用规范", level=1)
    doc.add_paragraph(
        f"本手册为央国企赴{jur_label}开展{archetypes_text}项目的开工前‘合规交底会’必备指引文件。"
        "严格落实国务院国资委《中央企业境外投资监督管理办法》（35号令）及《中央企业合规管理办法》，"
        "旨在贯通‘第一道防线（海外项目部）’、‘第二道防线（涉外法务合规官）’与‘第三道防线（集团内审纪检）’，"
        "实现海外基建全生命周期风险闭环留痕管理。"
    )

    # 3. 全生命周期五阶段实操合规指引 (SOP)
    doc.add_heading("二、 建设工程全生命周期五阶段合规实操指南 (SOP)", level=1)
    
    stages = [
        (
            "阶段 I：前期跟踪与立项尽调",
            "重点核查外资准入负面清单、BKPM/BOI 投资核准、国资委 35 号令重大投资前置论证、土地物权与习惯法 (Hak Ulayat) 征拆前置调查。",
            [
                "审查当地外资股比上限与负面清单限制（严禁违法代持）；",
                "核查合作伙伴世行/亚行除名黑名单；",
                "取得东道国省级以上政府土地征收与环评审批批文。"
            ]
        ),
        (
            "阶段 II：招投标与联营体组建",
            "属地法定施工牌照 (BUJKA/PCAB/CIDB) 合规、强制联营体 (KSO/JO) 协议连带责任隔离、反商业贿赂合规。",
            [
                "在印尼等法域依法组建 KSO 联营体，明确约定内部份额与对外连带责任；",
                "审查当地行会执照资质（PCAB v. Manila Water 虽放宽外资，公建项目仍受采购法特别约束）；",
                "落实全员海外反腐败合规廉洁承诺签署。"
            ]
        ),
        (
            "阶段 III：合同谈判与签署红线",
            "FIDIC 黄金原则偏离度审查、双语/三语文本同步签署与准据语言约定、见索即付保函反欺诈索兑限制条款、PPA/EPC 接口风险传导。",
            [
                "印尼项目严格依据 SEMA 3/2023 落实双语同步签署，杜绝合同无效抗辩风险；",
                "越南项目根据《2025新建设法》第86.2条重新设计预定损害赔偿 (LD) 封顶条款；",
                "马来西亚/新加坡项目坚决排除背靠背条款 (Pay-when-paid)；",
                "独立保函约定 SIAC/ICC 仲裁及紧急仲裁员 (EA) 救济机制。"
            ]
        ),
        (
            "阶段 IV：现场履约与过程留痕",
            "开工令与交地纪要签收、28天索赔通知标准化发函模板、外籍劳工配额与工作签台账、TKDN 本地采购认证留痕、工程款法定催告。",
            [
                "建立 28 天索赔通知前置预警机制，发生阻碍即刻发函保留权利；",
                "印尼新能源项目严格依能矿部 MEMR 11/2024 申请 TKDN 进口豁免许可；",
                "完整归档现场监理日记、地质异常情况与气象记录。"
            ]
        ),
        (
            "阶段 V：竣工移交、争议解决与退出清算",
            "缺陷责任期保函减额、CIPAA/SOPA 快速裁决应对、SIAC 紧急仲裁保函止付、中信保 30 日报损保全、《数据安全法》第36条跨境调证审批。",
            [
                "发生政府单方废约或索兑保函，30 日内向中信保提交正式《可能损失通知书》；",
                "向境外仲裁庭提交工程图纸前，依《数据安全法》第36条向我国主管机关完成前置审批；",
                "积极利用 CIPAA/SOPA 极速裁决机制追索被拖欠工程款。"
            ]
        )
    ]

    for stage_title, stage_desc, checklist in stages:
        doc.add_heading(stage_title, level=2)
        doc.add_paragraph(stage_desc)
        p_check = doc.add_paragraph()
        p_check.add_run("关键实操动作清单：\n").bold = True
        for c in checklist:
            p_check.add_run(f"  ☑ {c}\n")

    # 4. 国资委“三道防线”角色穿透审查表
    doc.add_heading("三、 国资委“三道防线”角色实操审查清单 (Checklists)", level=1)
    roles = [
        ("【第一道防线】海外市场部 / 项目现场经理审查要点", [
            "现场开工前确认交地纪要、红线图与规划许可证原件完整；",
            "严格监控分包商与设备商的 TKDN 认证供货批次凭证；",
            "任何由于业主原因导致的工期阻延，在 28 天内规范送达索赔意向通知函；",
            "严禁私自签署任何包含放弃索赔、推迟付款或减免业主责任的会议纪要。"
        ]),
        ("【第二道防线】涉外法务总监 / 专职合规官审查要点", [
            "合同文本双语对照复核，准据法与仲裁地优先选择新加坡 (SIAC) 或香港 (HKIAC)；",
            "对东道国强行法（如马来 CIPAA 35条、印尼 SEMA 3/2023）进行专用条件修正；",
            "核验中信保特别险投保状态与 30 日报损应急联络机制；",
            "审查跨境诉讼或调证材料，严守《数据安全法》第36条国家安全防线。"
        ]),
        ("【第三道防线】集团内部审计 / 纪检监察 / 国资穿透核查要点", [
            "核实境外项目重大资金流动与境外账户开立是否符合国资委 35 号令合规备案；",
            "穿透核查联营体与分包商是否存在境外利益输送或世行制裁连带风险；",
            "定期开展境外资产保全度与可追索性审计，核实母公司担保敞口风险。"
        ])
    ]
    for r_title, r_items in roles:
        doc.add_heading(r_title, level=2)
        p_r = doc.add_paragraph()
        for it in r_items:
            p_r.add_run(f"  • {it}\n")

    # 5. 示范反制与防御条款 (Model Counter-Clauses)
    doc.add_heading("四、 涉外工程合同示范防御条款 (Model Counter-Clauses)", level=1)
    doc.add_paragraph("以下中英双语示范条款可直接引入 FIDIC 专用条件 (Particular Conditions)，为央国企提供刚性合同保护：")

    counter_clauses = [
        (
            "示范条款 1：反习惯法征拆迟延免责条款 (Customary Land & Site Handover)",
            "The Employer shall deliver possession of the Site free from any encumbrances, including Hak Ulayat (customary land claims). Any delay in site handover shall entitle the Contractor to an Extension of Time (EOT) under Clause 8.4 and reimbursement of all associated Prolongation Costs.",
            "业主应移交不受任何权利负担限制的施工现场，包括原住民习惯法土地争议。任何现场移交迟延均应赋予承包商依据第8.4款主张工期顺延并全额索赔由此产生之窝工与延期成本的权利。"
        ),
        (
            "示范条款 2：见索即付保函索兑前置限制与止付约定 (Performance Security Stop-Call)",
            "The Employer shall give at least 21 days' prior written notice to the Contractor before making any demand under the Performance Security, specifying the alleged breach in detail. The Employer agrees that any wrongful call shall entitle the Contractor to immediate emergency arbitral injunction without proof of fraud.",
            "业主在向开证行发起履约保函索兑前，必须提前至少21天书面通知承包商并详述被指控之具体违约事实。业主在此确认并同意：任何无实质依据之索兑将赋予承包商直接申请国际仲裁紧急禁令止付的权利。"
        ),
        (
            "示范条款 3：东道国法律政策变更与电价传导条款 (Change in Law & Tariff Adjustment)",
            "If the Contractor suffers delay or incurs additional Cost as a result of any change in the Laws of the Country (including feed-in tariff adjustments or local content TKDN revisions), the Contract Price shall be adjusted accordingly and the Contractor shall be granted an Extension of Time.",
            "若承包商因东道国法律政策变更（包括光伏/电力上网电价下调或本地化率 TKDN 指标调整）遭受工期延误或额外成本，合同价金应作相应全额追加，且承包商应获准延长竣工工期。"
        )
    ]

    for c_title, c_en, c_zh in counter_clauses:
        doc.add_heading(c_title, level=2)
        p_c = doc.add_paragraph()
        p_c.add_run(f"【英文专用条件标准条文】\n{c_en}\n\n").italic = True
        p_c.add_run(f"【中文对照释义】\n{c_zh}")

    doc_stream = BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    return doc_stream
